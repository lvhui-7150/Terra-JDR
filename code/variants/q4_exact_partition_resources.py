#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""问题四第一小问：固定问题三任务后的精确分区与资源配置。"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Iterable, Iterator, Sequence

try:
    from openpyxl import load_workbook
except ImportError as exc:
    raise SystemExit("缺少 openpyxl。请运行：python -m pip install openpyxl") from exc

RESOURCE_TYPES = (
    "运输无人机-A", "运输无人机-B", "运输无人机-C",
    "共享电池-A", "共享电池-B", "共享电池-C",
    "中继无人机", "中继能源组件",
)
RESOURCE_LABELS = {
    "运输无人机-A": "A型运输无人机", "运输无人机-B": "B型运输无人机",
    "运输无人机-C": "C型运输无人机", "共享电池-A": "A型共享电池",
    "共享电池-B": "B型共享电池", "共享电池-C": "C型共享电池",
    "中继无人机": "中继无人机", "中继能源组件": "中继能源组件",
}
WORKLOAD_TYPES = (
    "服务区数", "货箱数", "运输架次数", "运输无人机占用_h",
    "运输能耗_kWh", "中继架次数", "中继服务_h", "中继能耗_kWh",
)
TIME_DIGITS = 6
BOX_PATTERN = re.compile(r"S\d{3}-[A-Z]{3}-\d{2,3}")
SERVICE_PATTERN = re.compile(r"S\d{3}")
TRANSPORT_ID_PATTERN = re.compile(r"Q3-T-\d{3,}")


@dataclass(frozen=True)
class TaskInterval:
    block_index: int
    resource: str
    start_s: float
    end_s: float
    task_id: str


@dataclass
class AtomicBlock:
    block_id: str
    service_areas: tuple[str, ...]
    transport_rows: list[dict[str, str]] = field(default_factory=list)
    relay_rows: list[dict[str, str]] = field(default_factory=list)
    workload: dict[str, float] = field(default_factory=dict)


@dataclass
class ProblemData:
    service_areas: tuple[str, ...]
    box_ids: tuple[str, ...]
    transport_rows: list[dict[str, str]]
    relay_rows: list[dict[str, str]]
    inventory: dict[str, int]
    blocks: list[AtomicBlock]
    intervals: list[TaskInterval]
    input_files: list[Path]
    q3_check_count: int


class DisjointSet:
    def __init__(self, items: Iterable[str]) -> None:
        self.parent = {item: item for item in items}
        self.rank = {item: 0 for item in items}

    def find(self, item: str) -> str:
        root = item
        while self.parent[root] != root:
            root = self.parent[root]
        while self.parent[item] != item:
            next_item = self.parent[item]
            self.parent[item] = root
            item = next_item
        return root

    def union(self, first: str, second: str) -> None:
        root_first, root_second = self.find(first), self.find(second)
        if root_first == root_second:
            return
        if self.rank[root_first] < self.rank[root_second]:
            root_first, root_second = root_second, root_first
        self.parent[root_second] = root_first
        if self.rank[root_first] == self.rank[root_second]:
            self.rank[root_first] += 1


def text(value: object) -> str:
    return "" if value is None else str(value).strip().lstrip("\ufeff")


def number(row: dict[str, str], names: Sequence[str], context: str) -> float:
    for name in names:
        if name in row and text(row[name]):
            try:
                value = float(row[name])
            except ValueError as exc:
                raise ValueError(f"{context}字段 {name} 不是数值：{row[name]!r}") from exc
            if not math.isfinite(value):
                raise ValueError(f"{context}字段 {name} 不是有限数")
            return value
    raise ValueError(f"{context}缺少字段：{' / '.join(names)}")


def string_field(row: dict[str, str], names: Sequence[str], context: str) -> str:
    for name in names:
        if name in row and text(row[name]):
            return text(row[name])
    raise ValueError(f"{context}缺少字段：{' / '.join(names)}")


def read_csv(path: Path) -> list[dict[str, str]]:
    for encoding in ("utf-8-sig", "utf-8"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                reader = csv.DictReader(handle)
                if not reader.fieldnames:
                    raise ValueError(f"CSV缺少表头：{path}")
                return [
                    {text(key): text(value) for key, value in row.items() if key}
                    for row in reader
                ]
        except UnicodeDecodeError:
            continue
    raise ValueError(f"不能用UTF-8读取CSV：{path}")


def find_csv(folder: Path, tokens: Sequence[str], excluded: Sequence[str] = ()) -> Path:
    matches = [
        path for path in folder.glob("*.csv")
        if all(token in path.name for token in tokens)
        and not any(token in path.name for token in excluded)
    ]
    if len(matches) != 1:
        found = ", ".join(path.name for path in matches) or "无匹配文件"
        raise ValueError(f"{folder} 中关键词 {tokens} 应唯一定位CSV，实际{len(matches)}个：{found}")
    return matches[0]


def find_header(ws, tokens: Sequence[str]) -> tuple[int, dict[str, int]]:
    for row_number, values in enumerate(ws.iter_rows(values_only=True), start=1):
        headers = {text(value): index for index, value in enumerate(values) if text(value)}
        if all(any(token in header for header in headers) for token in tokens):
            return row_number, headers
    raise ValueError(f"工作表 {ws.title} 找不到表头字段 {tokens}")


def column(headers: dict[str, int], token: str) -> int:
    matches = [index for header, index in headers.items() if token in header]
    if len(matches) != 1:
        raise ValueError(f"无法唯一定位表头 {token!r}：{list(headers)}")
    return matches[0]


def table(path: Path, sheet: str, tokens: Sequence[str]):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"{path.name} 缺少工作表 {sheet!r}")
        ws = wb[sheet]
        header_row, headers = find_header(ws, tokens)
        rows = list(ws.iter_rows(min_row=header_row + 1, values_only=True))
        return headers, rows
    finally:
        wb.close()


def load_reference_data(data_dir: Path):
    demand_path = data_dir / "物资需求与配送时限.xlsx"
    transport_path = data_dir / "运输无人机数据.xlsx"
    relay_path = data_dir / "中继无人机数据.xlsx"
    for path in (demand_path, transport_path, relay_path):
        if not path.is_file():
            raise FileNotFoundError(f"缺少题目附件：{path}")

    headers, rows = table(demand_path, "逐箱货箱清单", ("货箱编号", "服务区编号"))
    box_col, zone_col = column(headers, "货箱编号"), column(headers, "服务区编号")
    boxes, zones = [], set()
    for row in rows:
        box = text(row[box_col]) if box_col < len(row) else ""
        zone = text(row[zone_col]) if zone_col < len(row) else ""
        if not box and not zone:
            continue
        if not box or not zone or not BOX_PATTERN.fullmatch(box) or not SERVICE_PATTERN.fullmatch(zone):
            raise ValueError(f"货箱清单存在无效记录：{row}")
        boxes.append(box)
        zones.add(zone)
    if len(boxes) != 80 or len(set(boxes)) != 80 or len(zones) != 15:
        raise ValueError(f"附件清单校验失败：货箱记录/唯一数={len(boxes)}/{len(set(boxes))}，服务区数={len(zones)}")
    ordered_zones = tuple(sorted(zones, key=lambda item: int(item[1:])))

    headers, rows = table(transport_path, "数据", ("无人机编号", "机型编号"))
    id_col, type_col = column(headers, "无人机编号"), column(headers, "机型编号")
    drone_counts = {kind: 0 for kind in "ABC"}
    for row in rows:
        drone_id = text(row[id_col]) if id_col < len(row) else ""
        kind = text(row[type_col]) if type_col < len(row) else ""
        if drone_id.startswith("U") and kind in drone_counts:
            drone_counts[kind] += 1
    if drone_counts != {"A": 4, "B": 2, "C": 2}:
        raise ValueError(f"运输无人机库存与附件不符：{drone_counts}")

    headers, rows = table(transport_path, "数据", ("共享电池组总数", "等效完全充电时间"))
    type_col, count_col = column(headers, "机型编号"), column(headers, "共享电池组总数")
    batteries = {
        text(row[type_col]): int(row[count_col])
        for row in rows
        if type_col < len(row) and text(row[type_col]) in "ABC"
    }
    if batteries != {"A": 6, "B": 4, "C": 4}:
        raise ValueError(f"共享电池库存与附件不符：{batteries}")

    headers, rows = table(relay_path, "数据", ("中继无人机编号", "机型编号"))
    id_col, type_col = column(headers, "中继无人机编号"), column(headers, "机型编号")
    relay_drones = [
        text(row[id_col]) for row in rows
        if id_col < len(row) and type_col < len(row)
        and text(row[id_col]).startswith("R") and text(row[type_col]) == "R"
    ]
    if len(relay_drones) != 2:
        raise ValueError(f"中继无人机库存应为2架，实读{len(relay_drones)}架")

    headers, rows = table(relay_path, "数据", ("共享能源组件总数", "等效完全充电时间"))
    type_col, count_col = column(headers, "机型编号"), column(headers, "共享能源组件总数")
    components = [
        int(row[count_col]) for row in rows
        if type_col < len(row) and text(row[type_col]) == "R"
    ]
    if components != [6]:
        raise ValueError(f"中继能源组件库存应为6组，实读{components}")

    inventory = {
        "运输无人机-A": 4, "运输无人机-B": 2, "运输无人机-C": 2,
        "共享电池-A": 6, "共享电池-B": 4, "共享电池-C": 4,
        "中继无人机": 2, "中继能源组件": 6,
    }
    return ordered_zones, tuple(boxes), inventory


def validate_q3_checks(path: Path) -> int:
    rows = read_csv(path)
    if not rows or not all("检查项目" in row and "是否通过" in row for row in rows):
        raise ValueError("问题三约束检查表为空或缺少“检查项目/是否通过”字段")
    failures = [row for row in rows if text(row["是否通过"]) not in {"是", "通过", "PASS", "1", "TRUE"}]
    if failures:
        items = "；".join(f"{row['检查项目']}={row['是否通过']}" for row in failures)
        raise ValueError("问题三约束检查未全通过，拒绝作为输入：" + items)
    names = [text(row["检查项目"]) for row in rows]
    for token in ("货箱精确覆盖", "硬时限", "通信", "返航SOC"):
        if not any(token in name for name in names):
            raise ValueError(f"问题三约束检查缺少关键项目：{token}")
    return len(rows)


def parse_zones(row: dict[str, str], valid_zones: set[str], context: str) -> tuple[str, ...]:
    route = string_field(row, ("访问服务区顺序", "服务区顺序", "服务区集合"), context)
    zones = tuple(dict.fromkeys(SERVICE_PATTERN.findall(route)))
    invalid = [zone for zone in zones if zone not in valid_zones]
    if not zones or invalid:
        raise ValueError(f"{context}无法解析服务区：{route!r}；无效编号={invalid}")
    return zones


def parse_boxes(row: dict[str, str], context: str) -> tuple[str, ...]:
    value = string_field(row, ("货箱编号列表", "货箱清单", "货箱编号"), context)
    boxes = tuple(BOX_PATTERN.findall(value))
    if not boxes or len(boxes) != len(set(boxes)):
        raise ValueError(f"{context}货箱编号列表为空或有重复：{value!r}")
    return boxes


def parse_transport_rows(
    rows: list[dict[str, str]],
    service_areas: tuple[str, ...],
    box_ids: tuple[str, ...],
) -> list[dict[str, str]]:
    valid_zones, expected_boxes = set(service_areas), set(box_ids)
    if not rows:
        raise ValueError("问题三运输架次表为空")
    seen_tasks: set[str] = set()
    delivered: list[str] = []
    seen_zones: set[str] = set()
    for index, row in enumerate(rows, start=1):
        context = f"运输架次表第{index}行"
        task_id = string_field(row, ("架次编号", "运输架次编号"), context)
        if task_id in seen_tasks:
            raise ValueError(f"运输架次编号重复：{task_id}")
        seen_tasks.add(task_id)
        kind = string_field(row, ("机型编号",), context)
        if kind not in "ABC":
            raise ValueError(f"{context}机型编号无效：{kind}")
        zones = parse_zones(row, valid_zones, context)
        boxes = parse_boxes(row, context)
        declared_count = int(number(row, ("货箱数量",), context))
        if declared_count != len(boxes):
            raise ValueError(f"{context}货箱数量={declared_count}，但货箱编号数={len(boxes)}")
        start_s = number(row, ("开始时刻_s", "准备开始时刻_s"), context)
        takeoff_s = number(row, ("起飞时刻_s",), context)
        return_s = number(row, ("返回O01时刻_s", "返航时刻_s"), context)
        battery_ready_s = number(row, ("电池再次可用时刻_s", "电池再次可用时间_s"), context)
        energy = number(row, ("架次能耗_kWh", "运输能耗_kWh"), context)
        if not (0 <= start_s <= takeoff_s <= return_s <= battery_ready_s):
            raise ValueError(f"{context}时刻顺序错误：开始={start_s},起飞={takeoff_s},返航={return_s},电池可用={battery_ready_s}")
        if energy < 0:
            raise ValueError(f"{context}能耗不能为负")
        seen_zones.update(zones)
        delivered.extend(boxes)
        row.update({
            "__q4_task_id": task_id,
            "__q4_kind": kind,
            "__q4_zones": "、".join(zones),
            "__q4_box_ids": "、".join(boxes),
            "__q4_start": str(start_s),
            "__q4_takeoff": str(takeoff_s),
            "__q4_return": str(return_s),
            "__q4_battery_ready": str(battery_ready_s),
            "__q4_energy": str(energy),
        })
    duplicates = sorted({box_id for box_id in delivered if delivered.count(box_id) > 1})
    missing = sorted(expected_boxes - set(delivered))
    extras = sorted(set(delivered) - expected_boxes)
    if duplicates or missing or extras or len(delivered) != len(expected_boxes):
        raise ValueError(
            f"问题三货箱覆盖不符合附件：重复={duplicates[:8]}，漏送={missing[:8]}，"
            f"多出={extras[:8]}，总计={len(delivered)}箱"
        )
    if seen_zones != valid_zones:
        raise ValueError(f"问题三运输方案未覆盖15个服务区：{sorted(valid_zones - seen_zones)}")
    return rows


def parse_relay_rows(
    rows: list[dict[str, str]],
    transport_by_id: dict[str, dict[str, str]],
) -> list[dict[str, str]]:
    mission_ids: set[str] = set()
    for index, row in enumerate(rows, start=1):
        context = f"中继架次表第{index}行"
        mission_id = string_field(row, ("中继架次编号", "架次编号"), context)
        if mission_id in mission_ids:
            raise ValueError(
                f"中继架次编号重复：{mission_id}；请传入聚合后的物理中继架次表，不能传逐通信缺口记录。"
            )
        mission_ids.add(mission_id)
        support_text = string_field(row, ("服务运输架次", "保障运输架次", "运输架次编号"), context)
        support_ids = tuple(dict.fromkeys(TRANSPORT_ID_PATTERN.findall(support_text)))
        if not support_ids:
            raise ValueError(f"{context}无法解析保障的运输架次：{support_text!r}")
        missing = [task_id for task_id in support_ids if task_id not in transport_by_id]
        if missing:
            raise ValueError(f"{context}引用了不存在的运输架次：{missing}")
        launch_s = number(row, ("首次起飞时刻_s", "中继起飞时刻_s"), context)
        return_s = number(row, ("最后返航时刻_s", "返回O01时刻_s"), context)
        ready_s = number(row, ("能源组件最后可用时刻_s", "能源组件再次可用时刻_s"), context)
        energy = number(row, ("中继物理能耗_kWh", "中继能耗_kWh"), context)
        service_s = number(row, ("服务时间并集_s", "中继服务时长_s"), context)
        if not (0 <= launch_s <= return_s <= ready_s):
            raise ValueError(f"{context}时刻顺序错误：起飞={launch_s},返航={return_s},能源组件可用={ready_s}")
        if energy < 0 or service_s < 0:
            raise ValueError(f"{context}中继能耗或服务时长不能为负")
        row.update({
            "__q4_mission_id": mission_id,
            "__q4_supported_ids": "、".join(support_ids),
            "__q4_launch": str(launch_s),
            "__q4_return": str(return_s),
            "__q4_component_ready": str(ready_s),
            "__q4_energy": str(energy),
            "__q4_service": str(service_s),
        })
    return rows


def build_atomic_blocks(
    service_areas: tuple[str, ...],
    transport_rows: list[dict[str, str]],
    relay_rows: list[dict[str, str]],
) -> tuple[list[AtomicBlock], dict[str, int], dict[str, int]]:
    dsu = DisjointSet(service_areas)
    for row in transport_rows:
        zones = row["__q4_zones"].split("、")
        for zone in zones[1:]:
            dsu.union(zones[0], zone)
    transport_by_id = {row["__q4_task_id"]: row for row in transport_rows}
    for row in relay_rows:
        support_ids = row["__q4_supported_ids"].split("、")
        support_zones = [
            zone for task_id in support_ids
            for zone in transport_by_id[task_id]["__q4_zones"].split("、")
        ]
        for zone in support_zones[1:]:
            dsu.union(support_zones[0], zone)

    components: dict[str, list[str]] = {}
    for zone in service_areas:
        components.setdefault(dsu.find(zone), []).append(zone)
    ordered = sorted(
        (tuple(sorted(values, key=lambda zone: int(zone[1:]))) for values in components.values()),
        key=lambda values: int(values[0][1:]),
    )
    blocks = [
        AtomicBlock(
            block_id=f"C{index:03d}",
            service_areas=zones,
            workload={name: 0.0 for name in WORKLOAD_TYPES},
        )
        for index, zones in enumerate(ordered, start=1)
    ]
    zone_to_block = {
        zone: block_index
        for block_index, block in enumerate(blocks)
        for zone in block.service_areas
    }
    for block in blocks:
        block.workload["服务区数"] = float(len(block.service_areas))

    transport_to_block: dict[str, int] = {}
    for row in transport_rows:
        zones = row["__q4_zones"].split("、")
        indices = {zone_to_block[zone] for zone in zones}
        if len(indices) != 1:
            raise AssertionError(f"运输架次 {row['__q4_task_id']} 的服务区不在同一原子块")
        block_index = next(iter(indices))
        transport_to_block[row["__q4_task_id"]] = block_index
        block = blocks[block_index]
        block.transport_rows.append(row)
        block.workload["货箱数"] += len(row["__q4_box_ids"].split("、"))
        block.workload["运输架次数"] += 1
        block.workload["运输无人机占用_h"] += (float(row["__q4_return"]) - float(row["__q4_start"])) / 3600
        block.workload["运输能耗_kWh"] += float(row["__q4_energy"])

    relay_to_block: dict[str, int] = {}
    for row in relay_rows:
        indices = {transport_to_block[task_id] for task_id in row["__q4_supported_ids"].split("、")}
        if len(indices) != 1:
            raise AssertionError(f"中继架次 {row['__q4_mission_id']} 跨越多个原子块")
        block_index = next(iter(indices))
        relay_to_block[row["__q4_mission_id"]] = block_index
        block = blocks[block_index]
        block.relay_rows.append(row)
        block.workload["中继架次数"] += 1
        block.workload["中继服务_h"] += float(row["__q4_service"]) / 3600
        block.workload["中继能耗_kWh"] += float(row["__q4_energy"])
    return blocks, transport_to_block, relay_to_block


def build_intervals(
    blocks: list[AtomicBlock],
    transport_to_block: dict[str, int],
    relay_to_block: dict[str, int],
) -> list[TaskInterval]:
    intervals: list[TaskInterval] = []
    for block_index, block in enumerate(blocks):
        for row in block.transport_rows:
            task_id, kind = row["__q4_task_id"], row["__q4_kind"]
            start_s = float(row["__q4_start"])
            return_s = float(row["__q4_return"])
            ready_s = float(row["__q4_battery_ready"])
            intervals.append(TaskInterval(block_index, f"运输无人机-{kind}", start_s, return_s, task_id))
            intervals.append(TaskInterval(block_index, f"共享电池-{kind}", start_s, ready_s, task_id))
        for row in block.relay_rows:
            mission_id = row["__q4_mission_id"]
            launch_s, return_s = float(row["__q4_launch"]), float(row["__q4_return"])
            ready_s = float(row["__q4_component_ready"])
            intervals.append(TaskInterval(block_index, "中继无人机", launch_s, return_s, mission_id))
            intervals.append(TaskInterval(block_index, "中继能源组件", launch_s, ready_s, mission_id))
    return intervals


def load_problem_data(q3_dir: Path, data_dir: Path) -> ProblemData:
    if not q3_dir.is_dir():
        raise FileNotFoundError(f"问题三结果目录不存在：{q3_dir}")
    service_areas, box_ids, inventory = load_reference_data(data_dir)
    transport_path = find_csv(q3_dir, ("运输架次",))
    relay_path = find_csv(q3_dir, ("中继架次",), ("通信",))
    check_path = find_csv(q3_dir, ("约束检查",))
    check_count = validate_q3_checks(check_path)
    transport_rows = parse_transport_rows(read_csv(transport_path), service_areas, box_ids)
    transport_by_id = {row["__q4_task_id"]: row for row in transport_rows}
    relay_rows = parse_relay_rows(read_csv(relay_path), transport_by_id)
    blocks, transport_map, relay_map = build_atomic_blocks(service_areas, transport_rows, relay_rows)
    intervals = build_intervals(blocks, transport_map, relay_map)
    files = [
        transport_path, relay_path, check_path,
        data_dir / "物资需求与配送时限.xlsx",
        data_dir / "运输无人机数据.xlsx",
        data_dir / "中继无人机数据.xlsx",
    ]
    return ProblemData(
        service_areas, box_ids, transport_rows, relay_rows, inventory,
        blocks, intervals, files, check_count,
    )
