from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from functools import lru_cache
from itertools import combinations
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
import scipy.io as sio
from openpyxl import load_workbook


# ==============================
# 1. 文件位置与全局计算参数
# ==============================

ROOT = Path(__file__).resolve().parent
BASE_DATA = ROOT / "数据" / "无人机应急物资运输基础数据"
GEO_DATA = ROOT / "数据" / "镇龙乡地理空间数据" / "镇龙乡及周边地理数据"
DEM_FILE = GEO_DATA / "数字高程模型数据（DEM）" / "镇龙乡及周边30米DEM.mat"
OUTPUT_DIR = ROOT / "问题一第一问结果"
SUBMISSION_TEMPLATE = ROOT / "结果提交模板.xlsx"
SUBMISSION_OUTPUT = ROOT / "结果提交_问题一第一问.xlsx"

EARTH_RADIUS_M = 6_371_000.0
GRAVITY = 9.80665
JOULE_PER_KWH = 3.6e6
DEM_SAMPLE_STEP_M = 1.0  # 不超过30 m DEM像元边长的一半
BISECTION_ITERATIONS = 70
TOLERANCE = 1e-9


@dataclass(frozen=True)
class DroneType:
    """运输无人机机型参数。"""

    code: str
    name: str
    empty_mass_kg: float
    max_payload_kg: float
    volume_m3: float
    cruise_speed_mps: float
    empty_range_m: float
    full_range_m: float
    usable_energy_kwh: float
    reserve_ratio: float
    preparation_time_s: float
    loading_time_per_box_s: float
    handover_base_time_s: float
    handover_per_box_time_s: float
    climb_speed_mps: float
    descent_speed_mps: float
    climb_efficiency: float


@dataclass(frozen=True)
class Route:
    """调度中心到一个服务区的固定空间参数。"""

    service_id: str
    distance_m: float
    max_terrain_m: float
    cruise_altitude_m: float
    outbound_climb_m: float
    outbound_descent_m: float
    return_climb_m: float
    return_descent_m: float


# ==============================
# 2. 数据读取
# ==============================


def read_center_and_services() -> Tuple[dict, pd.DataFrame]:
    """读取调度中心和15个服务区的坐标、地面高程。"""

    file = BASE_DATA / "调度中心与服务区.xlsx"
    raw = pd.read_excel(file, sheet_name="数据", header=None)

    # 第3行是调度中心数据，第7行开始是服务区数据；这里显式按附件结构读取。
    center = {
        "id": str(raw.iloc[2, 0]),
        "longitude": float(raw.iloc[2, 2]),
        "latitude": float(raw.iloc[2, 3]),
        "ground_altitude_m": float(raw.iloc[2, 4]),
    }

    services = raw.iloc[6:21, [0, 2, 3, 4]].copy()
    services.columns = ["service_id", "longitude", "latitude", "ground_altitude_m"]
    services["service_id"] = services["service_id"].astype(str)
    for column in ["longitude", "latitude", "ground_altitude_m"]:
        services[column] = pd.to_numeric(services[column], errors="raise")

    if len(services) != 15:
        raise ValueError(f"服务区数量应为15，实际读取到 {len(services)} 个。")

    return center, services


def read_drone_types() -> List[DroneType]:
    """读取 A、B、C 三种运输无人机的参数。"""

    file = BASE_DATA / "运输无人机数据.xlsx"
    raw = pd.read_excel(file, sheet_name="数据", header=1)
    raw = raw[raw["机型编号"].isin(["A", "B", "C"]) & raw["最大载货质量（kg）"].notna()].copy()

    drones: List[DroneType] = []
    for _, row in raw.iterrows():
        reserve_ratio = float(row["返航电量下限（%）"]) / 100.0
        drones.append(
            DroneType(
                code=str(row["机型编号"]),
                name=str(row["机型名称"]),
                empty_mass_kg=float(row["含电池空载总质量（kg）"]),
                max_payload_kg=float(row["最大载货质量（kg）"]),
                volume_m3=float(row["可用装载体积（m³）"]),
                cruise_speed_mps=float(row["计划巡航速度（m/s）"]),
                empty_range_m=float(row["空载标准航程（m）"]),
                full_range_m=float(row["满载标准航程（m）"]),
                usable_energy_kwh=float(row["电池可用能量（kWh）"]),
                reserve_ratio=reserve_ratio,
                preparation_time_s=float(row["工位固定准备时间（s）"]),
                loading_time_per_box_s=float(row["每箱装载时间（s）"]),
                handover_base_time_s=float(row["接收点基础交接时间（s）"]),
                handover_per_box_time_s=float(row["每箱增加交接时间（s）"]),
                climb_speed_mps=float(row["最大爬升速度（m/s）"]),
                descent_speed_mps=float(row["最大下降速度（m/s）"]),
                climb_efficiency=float(row["爬升能耗效率"]),
            )
        )

    if {drone.code for drone in drones} != {"A", "B", "C"}:
        raise ValueError("运输无人机数据中未完整读取 A、B、C 三种机型。")

    return drones


def read_boxes() -> pd.DataFrame:
    """读取80个不可拆分货箱。"""

    file = BASE_DATA / "物资需求与配送时限.xlsx"
    boxes = pd.read_excel(file, sheet_name="逐箱货箱清单", header=0)
    boxes = boxes[
        [
            "货箱编号",
            "服务区编号",
            "物资类型",
            "单箱质量（kg）",
            "单箱体积（m³）",
        ]
    ].copy()
    boxes.columns = ["box_id", "service_id", "material_type", "mass_kg", "volume_m3"]
    boxes["box_id"] = boxes["box_id"].astype(str)
    boxes["service_id"] = boxes["service_id"].astype(str)
    boxes["mass_kg"] = pd.to_numeric(boxes["mass_kg"], errors="raise")
    boxes["volume_m3"] = pd.to_numeric(boxes["volume_m3"], errors="raise")

    if len(boxes) != 80:
        raise ValueError(f"货箱数量应为80，实际读取到 {len(boxes)} 个。")
    if boxes["box_id"].duplicated().any():
        raise ValueError("货箱编号存在重复，不能建立‘每箱恰好配送一次’约束。")

    return boxes


def load_dem() -> Tuple[np.ndarray, np.ndarray, np.ndarray]:
    """读取 DEM，并统一为经度递增、纬度递增的数组。"""

    dem = sio.loadmat(DEM_FILE)
    longitude = dem["longitude"].ravel().astype(float)
    latitude = dem["latitude"].ravel().astype(float)
    elevation = dem["dem"].astype(float)
    nodata = float(dem["nodata"].ravel()[0])
    elevation[elevation <= nodata] = np.nan

    # 原始纬度是递减的。翻转后，纬度轴和高程矩阵都变为递增方向。
    if latitude[0] > latitude[-1]:
        latitude = latitude[::-1]
        elevation = elevation[::-1, :]

    if not np.all(np.diff(longitude) > 0):
        raise ValueError("DEM经度坐标不是严格递增，无法进行采样。")
    if not np.all(np.diff(latitude) > 0):
        raise ValueError("DEM纬度坐标不是严格递增，无法进行采样。")

    return longitude, latitude, elevation


# ==============================
# 3. DEM采样与固定航段计算
# ==============================


def nearest_index(axis: np.ndarray, query: np.ndarray) -> np.ndarray:
    """返回每个查询值对应的最近栅格索引。"""

    right = np.searchsorted(axis, query, side="left")
    right = np.clip(right, 0, len(axis) - 1)
    left = np.clip(right - 1, 0, len(axis) - 1)

    choose_right = np.abs(axis[right] - query) < np.abs(axis[left] - query)
    return np.where(choose_right, right, left)


def sample_dem_nearest(
    longitude_query: np.ndarray,
    latitude_query: np.ndarray,
    longitude: np.ndarray,
    latitude: np.ndarray,
    elevation: np.ndarray,
) -> np.ndarray:
    """沿航段采样 DEM。使用最近栅格点，避免 NoData 双线性插值污染。"""

    if (
        longitude_query.min() < longitude.min()
        or longitude_query.max() > longitude.max()
        or latitude_query.min() < latitude.min()
        or latitude_query.max() > latitude.max()
    ):
        raise ValueError("航段超出 DEM 覆盖范围。")

    ix = nearest_index(longitude, longitude_query)
    iy = nearest_index(latitude, latitude_query)
    values = elevation[iy, ix]

    if np.isnan(values).any():
        raise ValueError("航段采样经过 DEM NoData 区域，请检查数据或改用插值方法。")
    return values


def local_xy(
    longitude: np.ndarray,
    latitude: np.ndarray,
    center_longitude: float,
    center_latitude: float,
) -> Tuple[np.ndarray, np.ndarray]:
    """把经纬度转换为以调度中心为原点的局部米制坐标。"""

    latitude_0 = np.radians(center_latitude)
    dx = EARTH_RADIUS_M * np.cos(latitude_0) * np.radians(longitude - center_longitude)
    dy = EARTH_RADIUS_M * np.radians(latitude - center_latitude)
    return dx, dy


def build_routes(
    center: dict,
    services: pd.DataFrame,
    longitude: np.ndarray,
    latitude: np.ndarray,
    elevation: np.ndarray,
) -> Dict[str, Route]:
    """根据题目固定的 O01 -> Si -> O01 形式计算所有服务区航段参数。"""

    route_dict: Dict[str, Route] = {}
    center_x, center_y = local_xy(
        np.array([center["longitude"]]),
        np.array([center["latitude"]]),
        center["longitude"],
        center["latitude"],
    )
    center_ground = center["ground_altitude_m"]

    for _, service in services.iterrows():
        service_id = service["service_id"]
        service_lon = float(service["longitude"])
        service_lat = float(service["latitude"])
        service_ground = float(service["ground_altitude_m"])

        service_x, service_y = local_xy(
            np.array([service_lon]),
            np.array([service_lat]),
            center["longitude"],
            center["latitude"],
        )
        distance_m = float(
            np.hypot(service_x[0] - center_x[0], service_y[0] - center_y[0])
        )

        sample_count = max(2, int(np.ceil(distance_m / DEM_SAMPLE_STEP_M)) + 1)
        ratio = np.linspace(0.0, 1.0, sample_count)
        profile_lon = center["longitude"] + ratio * (service_lon - center["longitude"])
        profile_lat = center["latitude"] + ratio * (service_lat - center["latitude"])
        profile_elevation = sample_dem_nearest(
            profile_lon, profile_lat, longitude, latitude, elevation
        )

        max_terrain_m = float(np.max(profile_elevation))
        cruise_altitude_m = max_terrain_m + 50.0

        # O01作业高度为地面高程，服务区作业高度为地面高程+30 m。
        center_work_altitude_m = center_ground
        service_work_altitude_m = service_ground + 30.0

        outbound_climb_m = cruise_altitude_m - center_work_altitude_m
        outbound_descent_m = cruise_altitude_m - service_work_altitude_m
        return_climb_m = cruise_altitude_m - service_work_altitude_m
        return_descent_m = cruise_altitude_m - center_work_altitude_m

        heights = [
            outbound_climb_m,
            outbound_descent_m,
            return_climb_m,
            return_descent_m,
        ]
        if min(heights) < -TOLERANCE:
            raise ValueError(f"{service_id} 的巡航高度低于节点作业高度，请检查 DEM 和节点高程。")

        route_dict[service_id] = Route(
            service_id=service_id,
            distance_m=distance_m,
            max_terrain_m=max_terrain_m,
            cruise_altitude_m=cruise_altitude_m,
            outbound_climb_m=outbound_climb_m,
            outbound_descent_m=outbound_descent_m,
            return_climb_m=return_climb_m,
            return_descent_m=return_descent_m,
        )

    return route_dict


# ==============================
# 4. 时间、能耗和安全载荷计算
# ==============================


def equivalent_range(drone: DroneType, payload_kg: float) -> float:
    """根据题目给出的载荷—航程关系计算等效航程。"""

    if not 0.0 <= payload_kg <= drone.max_payload_kg + TOLERANCE:
        raise ValueError(f"载荷 {payload_kg} kg 超过机型 {drone.code} 的质量范围。")

    ratio = np.clip(payload_kg / drone.max_payload_kg, 0.0, 1.0)
    return drone.empty_range_m - (drone.empty_range_m - drone.full_range_m) * ratio**1.5


def horizontal_energy(drone: DroneType, distance_m: float, payload_kg: float) -> float:
    """等效水平巡航能耗，单位 kWh。"""

    range_m = equivalent_range(drone, payload_kg)
    return drone.usable_energy_kwh * distance_m / range_m


def climb_energy(drone: DroneType, payload_kg: float, climb_height_m: float) -> float:
    """爬升附加能耗，单位 kWh；下降附加能耗按题目取0。"""

    total_mass_kg = drone.empty_mass_kg + payload_kg
    energy_joule = total_mass_kg * GRAVITY * climb_height_m / drone.climb_efficiency
    return energy_joule / JOULE_PER_KWH


def segment_energy(
    drone: DroneType,
    distance_m: float,
    climb_height_m: float,
    payload_kg: float,
) -> float:
    """计算单个航段总能耗。"""

    return horizontal_energy(drone, distance_m, payload_kg) + climb_energy(
        drone, payload_kg, climb_height_m
    )


def round_trip_energy(drone: DroneType, route: Route, payload_kg: float) -> float:
    """去程携带 payload，抵达后卸货，返程空载。"""

    outbound = segment_energy(
        drone, route.distance_m, route.outbound_climb_m, payload_kg
    )
    return_leg = segment_energy(drone, route.distance_m, route.return_climb_m, 0.0)
    return outbound + return_leg


def round_trip_flight_time(drone: DroneType, route: Route) -> float:
    """计算固定航段的去程+返程飞行时间，单位秒。"""

    outbound = (
        route.outbound_climb_m / drone.climb_speed_mps
        + route.distance_m / drone.cruise_speed_mps
        + route.outbound_descent_m / drone.descent_speed_mps
    )
    return_leg = (
        route.return_climb_m / drone.climb_speed_mps
        + route.distance_m / drone.cruise_speed_mps
        + route.return_descent_m / drone.descent_speed_mps
    )
    return outbound + return_leg


def group_operation_time(drone: DroneType, route: Route, box_count: int) -> float:
    """估算一个组批架次的累计作业时间，单位秒。"""

    return (
        round_trip_flight_time(drone, route)
        + drone.preparation_time_s
        + box_count * drone.loading_time_per_box_s
        + drone.handover_base_time_s
        + box_count * drone.handover_per_box_time_s
    )


def task_energy_limit(drone: DroneType) -> float:
    """返航安全余量约束下允许使用的最大整架次运输能耗。"""

    return (1.0 - drone.reserve_ratio) * drone.usable_energy_kwh


def return_remaining_energy(drone: DroneType, total_energy_kwh: float) -> float:
    """任务结束返航到中心后剩余的真实电量，单位 kWh。"""

    return drone.usable_energy_kwh - total_energy_kwh


def return_soc_ratio(drone: DroneType, total_energy_kwh: float) -> float:
    """任务结束返航到中心后的真实 SOC，取值为 0--1。"""

    return return_remaining_energy(drone, total_energy_kwh) / drone.usable_energy_kwh


def maximum_safe_payload(drone: DroneType, route: Route) -> Optional[float]:
    """用二分法求该服务区、该机型的最大安全载荷。"""

    budget = task_energy_limit(drone)

    # 如果空载往返都无法满足返航安全余量，则该机型在该服务区不可用。
    if round_trip_energy(drone, route, 0.0) > budget + TOLERANCE:
        return None

    low, high = 0.0, drone.max_payload_kg
    for _ in range(BISECTION_ITERATIONS):
        middle = (low + high) / 2.0
        if round_trip_energy(drone, route, middle) <= budget:
            low = middle
        else:
            high = middle
    return low


# ==============================
# 5. 货箱候选组批枚举
# ==============================


def enumerate_feasible_groups(
    service_id: str,
    boxes: pd.DataFrame,
    drones: Sequence[DroneType],
    route: Route,
    safe_payloads: Dict[str, Optional[float]],
) -> List[dict]:
    """枚举一个服务区内所有满足约束的“机型—货箱组批”候选。"""

    service_boxes = boxes[boxes["service_id"] == service_id].reset_index(drop=True)
    candidates: List[dict] = []

    # 货箱数量最多为15，完全枚举 2^15-1 个子集在本题规模下可行。
    for drone in drones:
        safe_payload = safe_payloads[drone.code]
        if safe_payload is None:
            continue

        for group_size in range(1, len(service_boxes) + 1):
            for selected_indices in combinations(range(len(service_boxes)), group_size):
                selected = service_boxes.iloc[list(selected_indices)]
                total_mass = float(selected["mass_kg"].sum())
                total_volume = float(selected["volume_m3"].sum())

                # 先检查质量和体积，尽早排除明显不可能的组合。
                if total_mass > safe_payload + TOLERANCE:
                    continue
                if total_volume > drone.volume_m3 + TOLERANCE:
                    continue

                total_energy = round_trip_energy(drone, route, total_mass)
                if total_energy > task_energy_limit(drone) + TOLERANCE:
                    continue

                remaining_energy = return_remaining_energy(drone, total_energy)
                return_soc = return_soc_ratio(drone, total_energy)
                if return_soc + TOLERANCE < drone.reserve_ratio:
                    continue

                box_ids = tuple(selected["box_id"].tolist())
                candidates.append(
                    {
                        "service_id": service_id,
                        "drone_type": drone.code,
                        "box_ids": box_ids,
                        "box_count": group_size,
                        "total_mass_kg": total_mass,
                        "total_volume_m3": total_volume,
                        "round_trip_energy_kwh": total_energy,
                        "task_energy_limit_kwh": task_energy_limit(drone),
                        "return_remaining_energy_kwh": remaining_energy,
                        "return_soc_ratio": return_soc,
                        "round_trip_flight_time_s": round_trip_flight_time(drone, route),
                        "estimated_operation_time_s": group_operation_time(
                            drone, route, group_size
                        ),
                    }
                )

    return candidates


# ==============================
# 6. 从候选组批中找一套完整可行方案
# ==============================


def find_one_feasible_cover(service_boxes: Sequence[str], candidates: List[dict]) -> List[dict]:
    """
    用深度优先搜索寻找一套“每个货箱恰好出现一次”的可行方案。

    这里不优化架次数、能耗或时间，只寻找可行解；第二问再在候选集合上做优化。
    搜索时优先尝试包含更多货箱的组批，通常能更快找到完整方案，但这不代表它是最优方案。
    """

    box_to_index = {box_id: index for index, box_id in enumerate(service_boxes)}
    all_mask = (1 << len(service_boxes)) - 1

    candidate_masks: List[int] = []
    for candidate in candidates:
        mask = 0
        for box_id in candidate["box_ids"]:
            mask |= 1 << box_to_index[box_id]
        candidate_masks.append(mask)

    # 每个货箱对应的候选组批索引，用于加速“优先处理候选最少的货箱”。
    candidates_by_box: Dict[int, List[int]] = {i: [] for i in range(len(service_boxes))}
    for candidate_index, mask in enumerate(candidate_masks):
        for box_index in range(len(service_boxes)):
            if mask & (1 << box_index):
                candidates_by_box[box_index].append(candidate_index)

    # 优先尝试货箱数多的组批；这是搜索顺序，不是第二问的优化目标。
    ordered_indices = sorted(
        range(len(candidates)),
        key=lambda index: (
            -candidates[index]["box_count"],
            candidates[index]["round_trip_energy_kwh"],
            candidates[index]["drone_type"],
            candidates[index]["box_ids"],
        ),
    )
    rank = {candidate_index: order for order, candidate_index in enumerate(ordered_indices)}

    dead_states = set()

    @lru_cache(maxsize=None)
    def search(covered_mask: int) -> Optional[Tuple[int, ...]]:
        if covered_mask == all_mask:
            return tuple()
        if covered_mask in dead_states:
            return None

        remaining_boxes = [
            box_index
            for box_index in range(len(service_boxes))
            if not (covered_mask & (1 << box_index))
        ]

        # 选择“当前可用候选组批最少”的未覆盖货箱，减少无效搜索。
        best_box = None
        best_options: List[int] = []
        for box_index in remaining_boxes:
            options = [
                candidate_index
                for candidate_index in candidates_by_box[box_index]
                if not (candidate_masks[candidate_index] & covered_mask)
            ]
            if not options:
                dead_states.add(covered_mask)
                return None
            if best_box is None or len(options) < len(best_options):
                best_box = box_index
                best_options = options

        best_options.sort(key=lambda index: rank[index])
        for candidate_index in best_options:
            candidate_mask = candidate_masks[candidate_index]
            result = search(covered_mask | candidate_mask)
            if result is not None:
                return (candidate_index,) + result

        dead_states.add(covered_mask)
        return None

    selected_indices = search(0)
    if selected_indices is None:
        return []
    return [candidates[index] for index in selected_indices]


# ==============================
# 7. 输出结果
# ==============================


def route_table(routes: Dict[str, Route]) -> pd.DataFrame:
    """把航段参数转换为表格。"""

    return pd.DataFrame(
        [
            {
                "服务区编号": route.service_id,
                "水平距离_m": route.distance_m,
                "沿线最高地面高程_m": route.max_terrain_m,
                "统一巡航海拔_m": route.cruise_altitude_m,
                "去程爬升高度_m": route.outbound_climb_m,
                "去程下降高度_m": route.outbound_descent_m,
                "返程爬升高度_m": route.return_climb_m,
                "返程下降高度_m": route.return_descent_m,
            }
            for route in routes.values()
        ]
    ).sort_values("服务区编号")


def safe_payload_table(
    routes: Dict[str, Route], drones: Sequence[DroneType]
) -> pd.DataFrame:
    """计算并输出15个服务区、3种机型的最大安全载荷。"""

    rows = []
    for service_id, route in routes.items():
        row = {"服务区编号": service_id}
        for drone in drones:
            safe_payload = maximum_safe_payload(drone, route)
            row[f"{drone.code}型最大安全载荷_kg"] = (
                np.nan if safe_payload is None else safe_payload
            )
            row[f"{drone.code}型空载往返能耗_kWh"] = round_trip_energy(drone, route, 0.0)
            row[f"{drone.code}型任务能量上限_kWh"] = task_energy_limit(drone)
        rows.append(row)
    return pd.DataFrame(rows).sort_values("服务区编号")


def write_submission_workbook(feasible_plan: Sequence[dict]) -> None:
    """从原始模板生成提交表副本，并写入第一小问的真实返航 SOC。"""

    workbook = load_workbook(SUBMISSION_TEMPLATE)
    worksheet = workbook["Q1_单点组批"]

    style_source = [worksheet.cell(2, column) for column in range(1, 10)]
    for row in worksheet.iter_rows(min_row=2, max_col=9):
        for cell in row:
            cell.value = None

    for row_number, candidate in enumerate(feasible_plan, start=2):
        values = [
            f"Q1-{row_number - 1:03d}",
            candidate["service_id"],
            candidate["drone_type"],
            "、".join(candidate["box_ids"]),
            candidate["total_mass_kg"],
            candidate["total_volume_m3"],
            candidate["estimated_operation_time_s"],
            candidate["round_trip_energy_kwh"],
            100.0 * candidate["return_soc_ratio"],
        ]
        for column, value in enumerate(values, start=1):
            source = style_source[column - 1]
            cell = worksheet.cell(row_number, column, value)
            cell.font = copy(source.font)
            cell.fill = copy(source.fill)
            cell.border = copy(source.border)
            cell.alignment = copy(source.alignment)
            cell.protection = copy(source.protection)
            cell.number_format = source.number_format

    for row_number in range(2, 2 + len(feasible_plan)):
        worksheet.cell(row_number, 5).number_format = "0.00"
        worksheet.cell(row_number, 6).number_format = "0.000"
        worksheet.cell(row_number, 7).number_format = "0.0"
        worksheet.cell(row_number, 8).number_format = "0.000"
        worksheet.cell(row_number, 9).number_format = "0.00"

    workbook.save(SUBMISSION_OUTPUT)


def main() -> None:
    print("开始读取问题一第一问所需数据……")
    center, services = read_center_and_services()
    drones = read_drone_types()
    boxes = read_boxes()
    longitude, latitude, elevation = load_dem()

    print(f"调度中心：{center['id']}")
    print(f"服务区数量：{len(services)}")
    print(f"货箱数量：{len(boxes)}")
    print(f"机型：{', '.join(drone.code for drone in drones)}")

    print("正在根据固定 O01 -> Si -> O01 航段提取 DEM 剖面……")
    routes = build_routes(center, services, longitude, latitude, elevation)

    print("正在计算各服务区、各机型的最大安全载荷……")
    safe_payloads_by_service: Dict[str, Dict[str, Optional[float]]] = {}
    for service_id, route in routes.items():
        safe_payloads_by_service[service_id] = {
            drone.code: maximum_safe_payload(drone, route) for drone in drones
        }

    print("正在枚举可行货箱组批……")
    all_candidates: List[dict] = []
    feasible_plan: List[dict] = []
    plan_check_rows: List[dict] = []
    drone_by_code = {drone.code: drone for drone in drones}

    for service_id in sorted(routes):
        route = routes[service_id]
        candidates = enumerate_feasible_groups(
            service_id,
            boxes,
            drones,
            route,
            safe_payloads_by_service[service_id],
        )
        all_candidates.extend(candidates)

        service_boxes = boxes.loc[boxes["service_id"] == service_id, "box_id"].tolist()
        one_plan = find_one_feasible_cover(service_boxes, candidates)
        feasible_plan.extend(one_plan)

        covered = [box_id for candidate in one_plan for box_id in candidate["box_ids"]]
        energy_violations = sum(
            candidate["round_trip_energy_kwh"]
            > candidate["task_energy_limit_kwh"] + TOLERANCE
            for candidate in one_plan
        )
        return_soc_violations = sum(
            candidate["return_soc_ratio"]
            + TOLERANCE
            < drone_by_code[candidate["drone_type"]].reserve_ratio
            for candidate in one_plan
        )
        if energy_violations or return_soc_violations:
            raise AssertionError(
                f"{service_id} 方案违反整架次能耗或返航SOC下限约束。"
            )

        plan_check_rows.append(
            {
                "服务区编号": service_id,
                "货箱总数": len(service_boxes),
                "可行候选组批数量": len(candidates),
                "可行方案架次数": len(one_plan),
                "货箱是否恰好覆盖一次": sorted(covered) == sorted(service_boxes)
                and len(covered) == len(set(covered)),
                "整架次能耗违规次数": energy_violations,
                "返航SOC违规次数": return_soc_violations,
                "最低返航SOC_pct": min(
                    (100.0 * candidate["return_soc_ratio"] for candidate in one_plan),
                    default=float("nan"),
                ),
            }
        )

        if not one_plan:
            print(f"警告：{service_id} 未找到覆盖全部货箱的可行方案。")

    print("正在写出结果文件……")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    route_table(routes).to_csv(OUTPUT_DIR / "航段参数.csv", index=False, encoding="utf-8-sig")
    safe_payload_table(routes, drones).to_csv(
        OUTPUT_DIR / "最大安全载荷.csv", index=False, encoding="utf-8-sig"
    )

    candidate_columns = [
        "服务区编号",
        "机型类型",
        "货箱编号列表",
        "货箱数量",
        "总质量_kg",
        "总体积_m3",
        "往返能耗_kWh",
        "任务能量上限_kWh",
        "返航剩余能量_kWh",
        "返航SOC_pct",
        "往返飞行时间_s",
        "估计架次作业时间_s",
    ]
    candidate_rows = []
    for candidate in all_candidates:
        candidate_rows.append(
            {
                "服务区编号": candidate["service_id"],
                "机型类型": candidate["drone_type"],
                "货箱编号列表": "、".join(candidate["box_ids"]),
                "货箱数量": candidate["box_count"],
                "总质量_kg": candidate["total_mass_kg"],
                "总体积_m3": candidate["total_volume_m3"],
                "往返能耗_kWh": candidate["round_trip_energy_kwh"],
                "任务能量上限_kWh": candidate["task_energy_limit_kwh"],
                "返航剩余能量_kWh": candidate["return_remaining_energy_kwh"],
                "返航SOC_pct": 100.0 * candidate["return_soc_ratio"],
                "往返飞行时间_s": candidate["round_trip_flight_time_s"],
                "估计架次作业时间_s": candidate["estimated_operation_time_s"],
            }
        )
    pd.DataFrame(candidate_rows, columns=candidate_columns).to_csv(
        OUTPUT_DIR / "可行候选组批.csv", index=False, encoding="utf-8-sig"
    )

    plan_rows = []
    for sortie_number, candidate in enumerate(feasible_plan, start=1):
        plan_rows.append(
            {
                "架次编号": f"Q1-{sortie_number:03d}",
                "服务区编号": candidate["service_id"],
                "机型类型": candidate["drone_type"],
                "货箱编号列表": "、".join(candidate["box_ids"]),
                "货箱数量": candidate["box_count"],
                "总质量_kg": candidate["total_mass_kg"],
                "总体积_m3": candidate["total_volume_m3"],
                "往返能耗_kWh": candidate["round_trip_energy_kwh"],
                "任务能量上限_kWh": candidate["task_energy_limit_kwh"],
                "返航剩余能量_kWh": candidate["return_remaining_energy_kwh"],
                "返航SOC_pct": 100.0 * candidate["return_soc_ratio"],
                "估计架次作业时间_s": candidate["estimated_operation_time_s"],
            }
        )
    pd.DataFrame(plan_rows).to_csv(
        OUTPUT_DIR / "第一问可行方案.csv", index=False, encoding="utf-8-sig"
    )

    pd.DataFrame(plan_check_rows).to_csv(
        OUTPUT_DIR / "第一问方案检查.csv", index=False, encoding="utf-8-sig"
    )

    write_submission_workbook(feasible_plan)

    minimum_return_soc = min(
        100.0 * candidate["return_soc_ratio"] for candidate in feasible_plan
    )

    print("\n计算完成。")
    print(f"输出目录：{OUTPUT_DIR}")
    print(f"提交表副本：{SUBMISSION_OUTPUT}")
    print(f"可行候选组批总数：{len(all_candidates)}")
    print(f"找到的可行方案架次数：{len(feasible_plan)}")
    print(f"最低返航完成后 SOC：{minimum_return_soc:.4f}%")
    print("注意：第一问可行方案只保证满足约束，不代表第二问意义下的最优方案。")


if __name__ == "__main__":
    main()
